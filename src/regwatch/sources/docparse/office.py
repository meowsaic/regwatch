"""Office 附件（.doc / .docx）正文提取。

中基协与证监会的决定书附件既有老式 OLE 复合文档（.doc），
也有 OOXML 包（.docx）；前者无官方解析器，这里做尽力而为的
WordDocument 流文本还原，并带 OLE 目录兜底。
"""

from __future__ import annotations

import re
from itertools import pairwise

from ...logging_setup import get_logger
from ..common import progress
from ..http import shared_session
from .pdf import _extract_text_from_pdf_bytes

logger = get_logger("sources.docparse")


def _detect_doc_format(content: bytes) -> str:
    """根据 magic bytes 检测文档实际格式。

    CSRC 附件常见三种情况：
    - OOXML (.docx)：PK 开头的 ZIP
    - OLE 复合文档 (.doc 或 .docx 误命名)：D0CF11E0 开头
    - PDF：%PDF 开头

    Args:
        content: 文件字节内容。

    Returns:
        'ole' | 'ooxml' | 'pdf' | 'unknown'
    """
    if content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "ole"
    if content[:4] in (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08"):
        return "ooxml"
    if content[:5] == b"%PDF-":
        return "pdf"
    return "unknown"


#: FIB 中 fcClx / lcbClx 的偏移（Word 97+ Binary Format）
_FIB_FC_CLX_OFFSET = 0x01A2
_FIB_LCB_CLX_OFFSET = 0x01A6


def _clx_from_fib(word_stream: bytes, table_stream: bytes) -> bytes:
    """按 FIB 规范读取 CLX（首选路径）。

    Word 97+ 规定 fcClx/lcbClx 位于 WordDocument 流偏移 0x01A2；
    盲扫 0x02 标志会命中 table 流里的假阳性（WPS 生成的 .doc 常见），
    进而把不相关字节当 piece table 解出成片乱码。
    """
    if len(word_stream) < _FIB_LCB_CLX_OFFSET + 4 or not table_stream:
        return b""
    fc_clx = int.from_bytes(word_stream[_FIB_FC_CLX_OFFSET : _FIB_FC_CLX_OFFSET + 4], "little")
    lcb_clx = int.from_bytes(word_stream[_FIB_LCB_CLX_OFFSET : _FIB_LCB_CLX_OFFSET + 4], "little")
    if 0 <= fc_clx < len(table_stream) and 0 < lcb_clx <= len(table_stream) - fc_clx:
        return table_stream[fc_clx : fc_clx + lcb_clx]
    return b""


def _clx_by_scan(table_stream: bytes) -> bytes:
    """兜底：盲扫 Pcdt 的 0x02 标志（FIB 不可用时）。"""
    idx = 0
    while idx < len(table_stream) - 6:
        if table_stream[idx] == 0x02:
            lcb = int.from_bytes(table_stream[idx + 1 : idx + 5], "little")
            if 0 < lcb < len(table_stream) - idx - 5 and (lcb - 4) % 12 == 0:
                return table_stream[idx:]
        idx += 1
    return b""


def _decode_piece_table(clx: bytes, word_stream: bytes) -> str:
    """解析 CLX（Prc* + Pcdt）中的 piece table 并按 CP 顺序拼接文本。

    结构异常（CP 不单调、字符数与流大小矛盾等）时返回空串——
    宁缺毋滥，避免把随机字节解成大片乱码。
    """
    idx = 0
    while idx < len(clx):
        flag = clx[idx]
        if flag == 0x01:  # Prc：跳过
            if idx + 3 > len(clx):
                return ""
            cb = int.from_bytes(clx[idx + 1 : idx + 3], "little")
            idx += 3 + cb
            continue
        if flag != 0x02:
            return ""
        if idx + 5 > len(clx):
            return ""
        lcb = int.from_bytes(clx[idx + 1 : idx + 5], "little")
        plc = clx[idx + 5 : idx + 5 + lcb]
        return _decode_pieces(plc, word_stream)
    return ""


def _decode_pieces(plc: bytes, word_stream: bytes) -> str:
    """PLCFpcd = (n+1)*CP(4字节) + n*PCD(8字节)，逐片解码 Unicode / ANSI 文本。"""
    if len(plc) < 4 or (len(plc) - 4) % 12 != 0:
        return ""
    n = (len(plc) - 4) // 12
    if n <= 0 or n > 4096:
        return ""
    cps = [int.from_bytes(plc[i * 4 : (i + 1) * 4], "little") for i in range(n + 1)]
    if cps[0] != 0 or any(b < a for a, b in pairwise(cps)):
        return ""
    # 每字符至少占 1 字节，总字符数不可能超过文档流长度
    if cps[-1] <= 0 or cps[-1] > len(word_stream):
        return ""
    pcds_start = (n + 1) * 4

    text_parts: list[str] = []
    for i in range(n):
        pcd = plc[pcds_start + i * 8 : pcds_start + (i + 1) * 8]
        if len(pcd) < 8:
            return ""
        fc_raw = int.from_bytes(pcd[2:6], "little")
        # bit30=1 表示 ANSI 压缩编码（fc 需除以 2），bit30=0 表示 Unicode
        is_unicode = not (fc_raw & 0x40000000)
        fc = fc_raw & 0x3FFFFFFF
        char_count = cps[i + 1] - cps[i]
        if is_unicode:
            raw = word_stream[fc : fc + char_count * 2]
            text_parts.append(raw.decode("utf-16-le", errors="ignore"))
        else:
            raw = word_stream[fc // 2 : fc // 2 + char_count]
            text_parts.append(raw.decode("cp936", errors="ignore"))
    return "".join(text_parts)


def _extract_text_from_ole(content: bytes) -> str:
    """从 OLE 复合文档（.doc 旧格式）提取正文文本。

    优先按 FIB 规范读取 fcClx/lcbClx 定位 piece table（:func:`_clx_from_fib`），
    失败退回旧版盲扫（:func:`_clx_by_scan`）；结构校验不过时使用流内文本兜底。

    Args:
        content: .doc 文件字节内容。

    Returns:
        提取的纯文本；失败返回空字符串。
    """
    import io

    import olefile

    try:
        ole = olefile.OleFileIO(io.BytesIO(content))
    except Exception as e:
        logger.warning(f"  [Word文档] OLE 打开失败: {e}")
        return ""

    try:
        if not ole.exists("WordDocument"):
            return ""
        word_stream = ole.openstream("WordDocument").read()

        # piece table 存放在 0Table 或 1Table 流
        table_name = "0Table" if ole.exists("0Table") else "1Table"
        table_stream = ole.openstream(table_name).read() if ole.exists(table_name) else b""

        for clx in (_clx_from_fib(word_stream, table_stream), _clx_by_scan(table_stream)):
            if not clx:
                continue
            text = _decode_piece_table(clx, word_stream)
            if len(text) > 50:
                return text

        fallback = _fallback_extract_ole_text(word_stream)
        if len(fallback) > 50 and not _looks_like_garbage(fallback):
            return fallback
        return ""
    finally:
        ole.close()


_CJK_CHAR = re.compile(r"[\u4e00-\u9fff]")


def _looks_like_garbage(text: str) -> bool:
    """粗判文本是否为二进制乱码：黑体取样中「汉字 + 可见 ASCII」占比过低。"""
    sample = text[:4000]
    if not sample:
        return True
    readable = len(_CJK_CHAR.findall(sample)) + sum(
        1 for ch in sample if ch.isascii() and ch.isprintable()
    )
    return readable / len(sample) < 0.5


def _fallback_extract_ole_text(stream: bytes) -> str:
    """OLE piece table 解析失败时，从 WordDocument 流直接提取文本片段。

    将整个流按 UTF-16-LE 解码后，用正则匹配连续的中英文可打印片段。
    质量不如 piece table 解析，但能兜底。
    """
    text = stream.decode("utf-16-le", errors="ignore")
    chunks = re.findall(
        r"[\u4e00-\u9fff\u3000-\u303f\uff00-\uffef\w\s，。；：、（）《》“”‘’！？·…—\-.\d()\[\]/%]+",
        text,
    )
    return "".join(chunks)


def _extract_text_from_ooxml(content: bytes) -> str:
    """从 OOXML (.docx) 提取文本。

    优先使用 python-docx（同时提取段落与表格）；失败时回退到直接
    解析 ZIP 内的 word/document.xml。
    """
    import io

    try:
        import docx  # python-docx

        document = docx.Document(io.BytesIO(content))
        paragraphs: list[str] = []
        for para in document.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)
        # 也提取表格中的文本（行政处罚决定书可能用表格列当事人信息）
        for table in document.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))
        return "\n".join(paragraphs)
    except Exception as e:
        logger.warning(f"  [Word文档] python-docx 失败，回退到 XML 解析: {e}")
        return _extract_text_from_ooxml_xml(content)


def _extract_text_from_ooxml_xml(content: bytes) -> str:
    """python-docx 失败时，直接从 ZIP 中解析 word/document.xml 提取文本。"""
    import io
    import xml.etree.ElementTree as ET
    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            doc_name = None
            for name in zf.namelist():
                if name.endswith("word/document.xml"):
                    doc_name = name
                    break
            if not doc_name:
                return ""
            with zf.open(doc_name) as f:
                tree = ET.parse(f)
        root = tree.getroot()
        w_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        paragraphs: list[str] = []
        for para in root.iter(f"{w_ns}p"):
            texts = [t.text for t in para.iter(f"{w_ns}t") if t.text]
            if texts:
                paragraphs.append("".join(texts))
        return "\n".join(paragraphs)
    except Exception as e:
        logger.warning(f"  [Word文档] XML 解析失败: {e}")
        return ""


def extract_text_from_docx(doc_url: str) -> str | None:
    """下载 Word 文档并提取纯文本，自动检测实际格式。

    CSRC 地方局附件常见三类情况：
    1. 标准 .docx（OOXML/ZIP）→ python-docx
    2. .doc 旧格式（OLE 复合文档）→ olefile 解析 piece table
    3. 扩展名为 .docx 但实际是 .doc（OLE）→ 同 2

    通过 magic bytes 检测实际格式后分发，避免依赖扩展名。

    Args:
        doc_url: Word 文档的 URL。

    Returns:
        提取的纯文本；失败或文本过短返回 None。
    """
    try:
        session = shared_session()
        resp = session.get(doc_url)
        resp.raise_for_status()
        content = resp.content
    except Exception as e:
        logger.warning(f"  [Word文档] 下载失败 {doc_url}: {e}")
        return None

    fmt = _detect_doc_format(content)
    if fmt == "ole":
        text = _extract_text_from_ole(content)
    elif fmt == "ooxml":
        text = _extract_text_from_ooxml(content)
    elif fmt == "pdf":
        # 极少数 Word 链接实际指向 PDF
        progress("  [Word文档] 实际为 PDF，转用 PDF 提取")
        text = _extract_text_from_pdf_bytes(content)
    else:
        logger.warning(f"  [Word文档] 未知格式（magic={content[:8].hex()}），尝试 OLE 回退")
        text = _extract_text_from_ole(content)

    if not text or len(text) <= 50:
        logger.warning(f"  [Word文档] 提取文本过短（{len(text or '')} 字符，格式={fmt}）")
        return None
    if _looks_like_garbage(text):
        logger.warning(f"  [Word文档] 提取结果为乱码（格式={fmt}，{len(text)} 字符），已丢弃")
        return None

    text = re.sub(r"\n{3,}", "\n\n", text)
    progress(f"  [Word文档] 提取成功（格式={fmt}），{len(text)} 字符")
    return text
