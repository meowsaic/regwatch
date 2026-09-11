import json, os, re
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SUMMARIES_DIR = r"E:\Desktop\codes\AMAC\summaries"
OUTPUT = r"E:\Desktop\codes\内控不健全案例清单.xlsx"

IC_KEYWORDS_VIOLATION_TYPE = ["内控缺失", "内控制度不健全", "内控不健全", "内控机制不健全", "内控混乱", "内控失效", "内控缺失"]
IC_KEYWORDS_SUMMARY = [
    "内控缺失", "内控不健全", "内控机制不健全", "内控制度不健全",
    "内控机制", "内控混乱", "内控失效",
    "内控制度", "合规风控负责人", "合规风控"
]
IC_AREAS = {
    "内控制度未建立/未落实": ["内控制度未建立", "未建立.*内控", "内控制度形同虚设", "内控机制不健全", "未建立.*有效.*内控", "内控制度.*不完善", "内控制度.*缺失", "内控.*混乱", "内控.*失效", "未有效执行内控", "内控缺失", "内控制度不健全"],
    "合规风控人员缺位": ["合规风控负责人.*缺位", "合规风控负责人.*空缺", "合规风控负责人.*未实际履职", "无.*合规.*负责人", "合规.*风控.*空缺", "无在职高级管理人员", "合规风控负责人.*离职", "无合规风控负责人", "合规风控.*不到位", "合规风控负责人.*未全面履职"],
    "基金材料保管不当": ["未妥善保管.*材料", "材料.*缺失", "未保存.*材料", "无法提供.*材料", "基金.*资料.*缺失", "未妥善保存", "材料.*丢失", "档案.*不完整", "未妥善保管.*资料"],
    "混同经营/独立性缺失": ["混同经营", "无独立办公场所", "与.*混同", "共用办公场所", "公章.*由.*管理", "独立性.*缺失"],
    "信息隔离/利益冲突": ["信息隔离", "利益冲突", "关联交易", "信息隔离墙"],
    "人员资质/场所违规": ["无基金从业资格", "挂名.*未实际履职", "人员.*相互.*调用", "场所.*违规", "不符合.*登记条件"],
    "投资运作内控缺失": ["投资决策由.*主导", "投资.*杠杆超标", "投资.*风险.*控制", "未尽勤勉尽责"],
    "适当性管理不到位": ["适当性管理不到位", "适当性.*不.*到位", "未.*履行.*适当性", "投资者.*适当性"],
    "信息披露违规": ["信息披露违规", "未.*披露.*信息", "信息披露.*不.*完整", "信息披露.*不.*真实"],
}

def is_ic_related(summary):
    vt = summary.get("violation_type", "")
    vs = summary.get("violation_summary", "")
    for kw in IC_KEYWORDS_VIOLATION_TYPE:
        if kw in vt:
            return True
    for kw in IC_KEYWORDS_SUMMARY:
        if kw in vs:
            return True
    return False

def analyze_ic_areas(summary):
    text = summary.get("violation_summary", "")
    vt = summary.get("violation_type", "")
    areas = []
    for area_name, patterns in IC_AREAS.items():
        for pattern in patterns:
            if re.search(pattern, text) or re.search(pattern, vt):
                areas.append(area_name)
                break
    return list(set(areas))

def extract_ic_description(summary):
    vs = summary.get("violation_summary", "")
    phrases = []
    ic_patterns = [
        r"内控[^，。；]*[，。；]?", r"内部控制[^，。；]*[，。；]?",
        r"内控制度[^，。；]*[，。；]?", r"合规风控[^，。；]*[，。；]?",
        r"风控[^，。；]*[，。；]?"
    ]
    for p in ic_patterns:
        matches = re.findall(p, vs)
        for m in matches:
            m = m.strip("，。；、 ")
            if len(m) > 5 and m not in phrases:
                phrases.append(m)
    if phrases:
        return "；".join(phrases[:5])
    return vs[:200] if vs else ""

all_files = []
for f in os.listdir(SUMMARIES_DIR):
    if f.endswith("_summary.json") and f != "_summary_index.json":
        all_files.append(os.path.join(SUMMARIES_DIR, f))

rows = []
for filepath in all_files:
    with open(filepath, "r", encoding="utf-8") as f:
        try:
            s = json.load(f)
        except json.JSONDecodeError:
            continue
    if s.get("entity_type") != "机构":
        continue
    date_str = s.get("date", "")
    if not date_str:
        continue
    try:
        d = datetime.strptime(date_str, "%Y-%m-%d")
    except ValueError:
        continue
    if d < datetime(2020, 1, 1):
        continue
    if not is_ic_related(s):
        continue
    areas = analyze_ic_areas(s)
    ic_desc = extract_ic_description(s)
    rows.append({
        "case_id": s.get("case_id", ""),
        "date": date_str,
        "entity": s.get("punished_entity", ""),
        "punishment": s.get("punishment", ""),
        "violation_type": s.get("violation_type", ""),
        "ic_areas": "、".join(areas) if areas else "",
        "ic_desc": ic_desc,
    })

rows.sort(key=lambda x: x["date"])

wb = Workbook()
ws = wb.active
ws.title = "内控不健全案例清单"

headers = ["案例编号", "处罚日期", "机构名称", "处罚类型", "违规类型", "内控问题环节", "内控问题具体描述"]
header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="4472C4")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

data_font = Font(name="Arial", size=10)
data_align = Alignment(vertical="top", wrap_text=True)
alt_fill = PatternFill("solid", fgColor="F2F7FB")

for i, row in enumerate(rows, 2):
    values = [row["case_id"], row["date"], row["entity"], row["punishment"], row["violation_type"], row["ic_areas"], row["ic_desc"]]
    for col, val in enumerate(values, 1):
        cell = ws.cell(row=i, column=col, value=val)
        cell.font = data_font
        cell.alignment = data_align
        cell.border = thin_border
        if i % 2 == 0:
            cell.fill = alt_fill

ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 35
ws.column_dimensions["D"].width = 18
ws.column_dimensions["E"].width = 35
ws.column_dimensions["F"].width = 28
ws.column_dimensions["G"].width = 60

ws.auto_filter.ref = f"A1:G{len(rows)+1}"
ws.freeze_panes = "A2"

summary_row = len(rows) + 3
ws.cell(row=summary_row, column=1, value="统计说明").font = Font(name="Arial", bold=True, size=11)
ws.cell(row=summary_row+1, column=1, value=f"统计范围: 2020-01-01至今，仅机构受处罚案例").font = data_font
ws.cell(row=summary_row+2, column=1, value=f"内控相关案例总数: {len(rows)}件").font = data_font
ws.cell(row=summary_row+3, column=1, value=f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").font = data_font

wb.save(OUTPUT)
print(f"已导出: {OUTPUT} ({len(rows)}条记录)")
