# -*- coding: utf-8 -*-
"""完善「华宝投资优秀岗位创新成果奖申报-风控合规部」版申报表：
1. 成果简介具体化四份专题研究成果名称；
2. 详细内容·应用效果（3）决策支撑处，补充专题报告的具体内容与结论；
3. 填充经济效益（计算依据/计算方式/年净增）；
4. 填充知识产权汇总表（软件著作权）及基本情况「其他」栏。
在原文件基础上生成“-已完善”版本，不覆盖原件。
"""
import copy
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

SRC = r"E:\Desktop\codes\华宝投资优秀岗位创新成果奖申报-风控合规部-20260911.docx"
DST = r"E:\Desktop\codes\华宝投资优秀岗位创新成果奖申报-风控合规部-20260911-已完善.docx"

doc = Document(SRC)

def _first_rpr(para):
    for r in para.runs:
        rr = r._element.find(qn("w:rPr"))
        if rr is not None:
            return copy.deepcopy(rr)
    return None

def set_para_text(para, text, ref_para=None):
    """重写段落文本，保留原 pPr、run 用参考段落字体。"""
    ppr = para._p.find(qn("w:pPr"))
    rpr = _first_rpr(para)
    if rpr is None and ref_para is not None:
        rpr = _first_rpr(ref_para)
    for child in list(para._p):
        para._p.remove(child)
    if ppr is not None:
        para._p.append(ppr)
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    para._p.append(r)

def insert_after(para, text, ref_para=None):
    """在 para 之后插入新段落，继承参考段落的 pPr 与字体。"""
    new_p = OxmlElement("w:p")
    ref = ref_para if ref_para is not None else para
    ref_ppr = ref._p.find(qn("w:pPr"))
    if ref_ppr is not None:
        new_p.append(copy.deepcopy(ref_ppr))
    rpr = _first_rpr(ref)
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    new_p.append(r)
    para._p.addnext(new_p)
    return Paragraph(new_p, para._parent)

def set_cell(cell, text):
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    first = cell.paragraphs[0]
    set_para_text(first, text, ref_para=first)

# ============================================================
# 1. 成果简介：具体化专题报告名称
# ============================================================
t1 = doc.tables[1].cell(0, 0)
jianjie = [
    t1.paragraphs[0].text,
    ("平台覆盖全市场数千条监管处分案例，提供网页看板与命令行双入口，支持按机构类型、违规类型、"
     "来源局、日期、关键词等多维度组合检索，可一键生成违规分布、处罚构成、机构与人员对比、法规引用、"
     "时间趋势等统计报告，并已在此基础之上形成《私募基金管理人非专业化运营边界分析报告》《AMAC内控不健全"
     "专项分析报告》《CSRC内控不健全专项分析报告》《私募基金管理人高级管理人员违规兼职典型案例分析》等"
     "专题研究成果。成果已应用于公司重大事项的前置合规评估，为经营管理决策提供专业数据支撑，促进科学决策与合规经营。"),
]
set_para_text(t1.paragraphs[1], jianjie[1], ref_para=t1.paragraphs[0])

# ============================================================
# 2. 详细内容·应用效果（3）后补充专题成果详情
# ============================================================
t2 = doc.tables[2].cell(0, 0)
paras = t2.paragraphs
# 定位“（3）决策支撑”段落
idx3 = None
for i, p in enumerate(paras):
    if p.text.startswith("（3）决策支撑"):
        idx3 = i
        break
assert idx3 is not None, "未找到应用效果（3）决策支撑段落"

detail = ("截至目前，平台已支撑形成四份专题研究成果：《私募基金管理人非专业化运营边界分析报告》基于"
          "2023年1月至2026年5月125份AMAC纪律处分案例，将“非专业化运营”归纳为兼营债券发行承销、财务顾问、"
          "借贷放贷、居间中介、定向融资、通道业务、代持理财等十余类典型行为并厘清认定逻辑；《AMAC内控不健全"
          "专项分析报告》基于2020年以来554个机构处分案例，揭示内控缺失占比由2021年的20%升至2025年的24.5%"
          "及高发环节分布；《CSRC内控不健全专项分析报告》基于2022年以来2019个案例，显示内控缺失占比由2022年的"
          "16.0%升至2026年的27.6%，印证监管对内控要求的持续收紧；《私募基金管理人高级管理人员违规兼职典型案例"
          "分析》覆盖两体系典型案例，得出“违规兼职通常与内控缺失、登记信息失实等多项违规一并处罚”的判断，"
          "为兼职合规审查提供口径参考。")
insert_after(paras[idx3], detail)

# ============================================================
# 3. 经济效益：填充三个标题下正文
# ============================================================
t3 = doc.tables[3].cell(0, 0)
p3 = t3.paragraphs
ref_body = t2.paragraphs[1]   # 用详细内容正文段作为格式参照

jj_yj = ("以平台替代法务合规人员人工“检索—阅读—归类—统计”监管处分案例所节约的工作量为测算口径，"
         "将成果实施前后的合规分析工时差，结合人力成本折算为年净增经济效益。")
jj_fs = ("（1）单次作业节约：一次全口径监管案例检索与统计，人工约需4人·天，平台上线后约需0.5人·天"
         "（主要用于结果核对），单次节约约3.5人·天。（2）年度作业频次：按年度开展重大事项前置合规评估、"
         "季度及专题监管分析合计约20次测算，年节约工时约70人·天。（3）年净增经济效益＝年节约工时×"
         "法务人员日均人力成本（含薪酬、社保及管理费分摊）。")
jj_nd = ("按法务人员日均人力成本【待填：元/人·天】计算，年净增经济效益约为【待填：万元】。"
         "以上为测算口径，最终金额以推荐单位财务部门审核认定为准。")

# 空段落索引：1（计算依据）、3（计算方式）、5（年净增）
set_para_text(p3[1], jj_yj, ref_para=ref_body)
set_para_text(p3[3], jj_fs, ref_para=ref_body)
set_para_text(p3[5], jj_nd, ref_para=ref_body)

# ============================================================
# 4. 知识产权：基本情况「其他」栏 + 知识产权汇总表首行
# ============================================================
t0 = doc.tables[0]
set_cell(t0.rows[11].cells[4], "软件著作权（拟申请，详见《知识产权情况汇总表》）")

t7 = doc.tables[7]
r1 = t7.rows[1]
set_cell(r1.cells[0], "软件著作权")
set_cell(r1.cells[1], "华宝股权私募基金监管案例全口径智能采集与合规分析平台（软件）")
set_cell(r1.cells[2], "【待填：登记号/受理号】")
set_cell(r1.cells[3], "【待填：认定时间】")
set_cell(r1.cells[4], "如尚未登记请注明“拟申请/受理中”")

doc.save(DST)
print("已生成:", DST)